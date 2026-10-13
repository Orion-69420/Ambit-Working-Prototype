"""End-to-end Step 10 flow on a COPY of the real database: review, approve, export, change.

Skipped when the pipeline has not been run (no data/copilot.sqlite)."""
import shutil
from decimal import Decimal

import pymupdf
import pytest
from openpyxl import load_workbook

from copilot import config, db, export, report, review, verify, xbrl_extract

REAL = db.DEFAULT_DB
pytestmark = pytest.mark.skipif(not REAL.exists(), reason="pipeline not run: no database")
WHO = "Test reviewer"


@pytest.fixture
def env(tmp_path):
    path = tmp_path / "copy.sqlite"
    shutil.copy(REAL, path)
    return config.load_all(), db.connect(path), tmp_path


def review_all(conn, cfg, company, quarter, comment="checked against source"):
    rep = report.build(conn, cfg, company, quarter)
    for item in rep["items"]:
        review.record(conn, cfg, company, quarter, item, WHO, "approve", comment)


def test_cannot_approve_until_everything_reviewed(env):
    cfg, conn, _ = env
    with pytest.raises(report.ApprovalError, match="not yet reviewed"):
        report.approve(conn, cfg, "MM", "Q1FY27", WHO)


def test_full_flow_review_approve_export(env):
    cfg, conn, out = env
    review_all(conn, cfg, "MM", "Q1FY27")
    rep = report.approve(conn, cfg, "MM", "Q1FY27", WHO)
    assert report.status(conn, cfg, "MM", "Q1FY27")["status"] == "approved"

    snap, files = export.export(conn, cfg, "MM", "Q1FY27", out_dir=out)
    assert snap["status"] == "approved" and snap["approved_by"] == WHO
    assert [f.name for f in files] == ["MM_Q1FY27_APPROVED.xlsx", "MM_Q1FY27_APPROVED.pdf"]

    # Excel and PDF carry the same figures, from the same snapshot
    wb = load_workbook(files[0])
    xl_text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    pdf_text = " ".join(p.get_text() for p in pymupdf.open(files[1]))
    pdf_flat = " ".join(pdf_text.split())
    sec = export.sections(snap)
    for row in sec["metrics"]:
        assert row[1] in xl_text
        assert " ".join(row[1].replace("₹", "Rs ").split()) in pdf_flat
    assert "Audit log" in wb.sheetnames
    assert any(r[2] == "report_approved" for r in wb["Audit log"].iter_rows(values_only=True))
    # exporting does not approve, and is audited
    assert conn.execute("SELECT count(*) FROM audit_log WHERE action = 'report_exported'").fetchone()[0] >= 1


def test_draft_export_is_labelled(env):
    cfg, conn, out = env
    snap, files = export.export(conn, cfg, "INFY", "Q1FY27", out_dir=out)
    assert snap["status"] == "draft" and files[0].name.endswith("_DRAFT.xlsx")
    assert load_workbook(files[0])["Report"]["D1"].value == "DRAFT"
    assert "DRAFT" in pymupdf.open(files[1])[0].get_text()


def test_change_after_approval_returns_to_draft(env):
    cfg, conn, _ = env
    review_all(conn, cfg, "MM", "Q1FY27")
    report.approve(conn, cfg, "MM", "Q1FY27", WHO)
    review.record(conn, cfg, "MM", "Q1FY27", "metric:mm_suv_volumes:", WHO, "correct",
                  "deck shows 174.8 on p30 after zoom", "174.8")
    st = report.status(conn, cfg, "MM", "Q1FY27")
    assert st["status"] == "draft" and st.get("returned")
    # the original extraction is kept beside the correction
    rows = conn.execute("SELECT value, method, supersedes, reason FROM figures WHERE company = 'MM'"
                        " AND metric = 'mm_suv_volumes' AND period = 'Q1FY27' ORDER BY id").fetchall()
    assert any(r["method"] == "pdf" for r in rows) and any(r["supersedes"] and r["reason"] for r in rows)
    row = next(r for r in report.build(conn, cfg, "MM", "Q1FY27")["content"]["rows"] if r["metric"] == "mm_suv_volumes")
    assert Decimal(row["current"]["value"]) == Decimal("174.8")


def test_low_confidence_needs_a_check_note_and_survives_rerun(env):
    cfg, conn, _ = env
    key = "metric:revenue:"
    with pytest.raises(review.ReviewError, match="low-confidence"):
        review.record(conn, cfg, "TITAN", "Q1FY27", key, WHO, "approve")
    review.record(conn, cfg, "TITAN", "Q1FY27", key, WHO, "approve", "sale of products + other operating revenue = 21,356")
    st = conn.execute("SELECT status FROM figures WHERE company = 'TITAN' AND metric = 'revenue' AND period = 'Q1FY27'"
                      " AND method = 'xbrl'").fetchone()[0]
    assert st == "manually_verified"
    # re-running extraction and checks keeps the analyst's check while the value is unchanged
    xbrl_extract.run(conn, cfg, "TITAN", "Q1FY27")
    verify.run(conn, cfg, "TITAN", "Q1FY27")
    st = conn.execute("SELECT status FROM figures WHERE company = 'TITAN' AND metric = 'revenue' AND period = 'Q1FY27'"
                      " AND method = 'xbrl'").fetchone()[0]
    assert st == "manually_verified"


def test_reviews_survive_a_rerun_of_the_pipeline(env):
    cfg, conn, _ = env
    review_all(conn, cfg, "MM", "Q1FY27")
    before = report.build(conn, cfg, "MM", "Q1FY27")
    from copilot import commentary, compare
    xbrl_extract.run(conn, cfg, "MM", "Q1FY27")
    verify.run(conn, cfg, "MM", "Q1FY27")
    compare.run(conn, cfg, "MM", "Q1FY27")
    commentary.run(conn, cfg, "MM", "Q1FY27")
    after = report.build(conn, cfg, "MM", "Q1FY27")
    assert after["fingerprint"] == before["fingerprint"]
    assert not report.blockers(conn, after)


def test_headline_uses_template_and_verified_figures_only(env):
    cfg, conn, _ = env
    c = report.build(conn, cfg, "TITAN", "Q1FY27")["content"]
    keys = {h["key"] for h in c["headline"]}
    assert "metric:revenue:" not in keys            # low confidence -> not in the headline
    assert 0 < len(c["headline"]) <= 5
    assert all(" the estimate of " in h["text"] and h["text"].endswith(")." ) for h in c["headline"])


@pytest.mark.parametrize("value, unit, text", [
    ("58187.57", "inr_crore", "₹58,188 crore"), ("311850", "inr_crore", "₹3,11,850 crore"),
    ("48.8", "inr_per_share", "₹48.80"), ("17.5252", "percent", "17.5%"), (None, "percent", "—")])
def test_value_formatting(value, unit, text):
    assert report.fmt_value(value, unit) == text
