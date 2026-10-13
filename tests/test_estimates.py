from decimal import Decimal

import pytest
from openpyxl import Workbook

from copilot import config, estimates, ingest


@pytest.fixture
def cfg():
    return config.load_all()


@pytest.fixture
def model(tmp_path):
    """A small model: a typed value, a float, a formula with no cached value, and an empty cell."""
    wb = Workbook()
    ws = wb.active
    ws.title = "MM"
    ws["D4"], ws["E4"] = "2026-04-10", "2026-07-10"
    ws["D5"], ws["E5"] = 62479, 57328.55
    ws["D6"] = "=D5*0.16"          # openpyxl stores formulas without a cached value
    path = tmp_path / "model.xlsx"
    wb.save(path)
    mmap = {"workbook": "model.xlsx", "sample": True, "companies": {"MM": {
        "sheet": "MM", "estimate_date": {"Q4FY26": "D4", "Q1FY27": "E4"},
        "cells": [{"metric": "revenue", "segment": None, "unit": "inr_crore", "Q4FY26": "D5", "Q1FY27": "E5"},
                  {"metric": "ebitda", "segment": None, "unit": "inr_crore", "Q4FY26": "D6", "Q1FY27": "E6"}]}}}
    return path, mmap


def test_reads_cached_values_with_cell_references(model):
    path, mmap = model
    ests, before, after = estimates.read(path, mmap)
    by = {(e["metric"], e["quarter"]): e for e in ests}
    assert by[("revenue", "Q4FY26")]["value"] == Decimal("62479")
    assert by[("revenue", "Q1FY27")]["value"] == Decimal("57328.55")      # float read back exactly
    assert by[("revenue", "Q1FY27")]["cell"] == "E5" and by[("revenue", "Q1FY27")]["sheet"] == "MM"
    assert by[("revenue", "Q4FY26")]["estimate_date"] == "2026-04-10"
    assert all(e["is_sample"] for e in ests)


def test_formula_without_cached_value_is_missing(model):
    path, mmap = model
    ests, _, _ = estimates.read(path, mmap)
    by = {(e["metric"], e["quarter"]): e for e in ests}
    assert by[("ebitda", "Q4FY26")]["value"] is None       # formula never evaluated
    assert by[("ebitda", "Q1FY27")]["value"] is None       # empty cell


def test_workbook_hash_unchanged_after_read(model):
    path, mmap = model
    sha = ingest.sha256(path)
    _, before, after = estimates.read(path, mmap)
    assert before == after == sha == ingest.sha256(path)


def test_map_validation_rejects_unknown_metric_and_unit(cfg, tmp_path):
    (tmp_path / "model_map.yaml").write_text(
        "workbook: x.xlsx\ncompanies:\n  MM:\n    sheet: MM\n    cells:\n"
        "    - {metric: sales, unit: inr_crore, Q4FY26: D5}\n"
        "    - {metric: revenue, unit: percent, Q4FY26: D6}\n", encoding="utf8")
    with pytest.raises(config.ConfigError) as e:
        estimates.load_map(cfg, tmp_path)
    assert "unknown metric 'sales'" in str(e.value) and "unit percent != inr_crore" in str(e.value)


def test_real_model_map_is_valid(cfg):
    m = estimates.load_map(cfg)
    assert m["sample"] is True and set(m["companies"]) == set(cfg.companies)


@pytest.mark.parametrize("raw, expected", [(24, Decimal("24")), (0.1, Decimal("0.1")), ("1,240.5", Decimal("1240.5")),
                                           (None, None), (True, None), ("n/a", None)])
def test_to_decimal(raw, expected):
    assert estimates.to_decimal(raw) == expected
