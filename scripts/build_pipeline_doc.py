"""Build docs/Pipeline_Documentation.docx in the house format of the Phase 1 procedure.

    python scripts/build_pipeline_doc.py

Uses docs/templates/ambit_style.docx (the procedure's styles and numbering, no content) and
fills it with the pipeline documentation. Counts are read from the database at build time so
the document matches the current state; run after `python run.py run` and the exports.

docs/Pipeline_Documentation.pdf is the same document exported from Word (File > Save as PDF, after
right-click > Update field on the table of contents), so it can be read on GitHub.
"""
import copy
import csv
import datetime as dt
import pathlib
import re
import sqlite3
import sys

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from copilot import config, db  # noqa: E402

TEMPLATE = ROOT / "docs" / "templates" / "ambit_style.docx"
OUT = ROOT / "docs" / "Pipeline_Documentation.docx"
REPO = "https://github.com/Orion-69420/Ambit-Working-Prototype"
MAROON, GRID, CELL_SIZE = "6B1A0E", "BFBFBF", "19"     # as in the procedure document
DECIMAL_NUM, BULLET_NUM = "1001", "1002"


class Doc:
    """Writes paragraphs and tables with the procedure document's formatting."""

    def __init__(self):
        self.d = docx.Document(TEMPLATE)
        self.body = self.d.element.body
        self.sect = self.body.sectPr
        self.numbering = self.d.part.numbering_part.element
        self.next_num = 3000
        self._sections = 0

    def _add(self, el):
        self.sect.addprevious(el)
        return el

    def _p(self, style, text="", color=None, runs=None):
        p = self.d.add_paragraph()           # appended before sectPr by python-docx
        if style:
            p.style = self.d.styles[style]
        for chunk, bold in (runs or self._inline(text)):
            r = p.add_run(chunk)
            r.bold = bold or None
            if color:
                r._r.get_or_add_rPr().append(parse_xml(f'<w:color {nsdecls("w")} w:val="{color}"/>'))
        return p

    @staticmethod
    def _inline(text):
        """'**bold** rest' -> [(text, bold)]"""
        parts = re.split(r"(\*\*[^*]+\*\*)", text)
        return [(s[2:-2], True) if s.startswith("**") else (s, False) for s in parts if s]

    def title(self, title, subtitle, author, date):
        self._p("Title", title)
        self._p("Subtitle", subtitle)
        self._p("Author", author)
        self._p("Date", date)
        self._add(parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
            '<w:docPartUnique/></w:docPartObj></w:sdtPr><w:sdtContent>'
            '<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>Table of Contents</w:t></w:r></w:p>'
            '<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/>'
            '<w:instrText xml:space="preserve">TOC \\o "1-1" \\h \\z \\u</w:instrText>'
            '<w:fldChar w:fldCharType="separate"/><w:fldChar w:fldCharType="end"/></w:r></w:p>'
            '</w:sdtContent></w:sdt>'))

    def h1(self, text):
        if self._sections:                    # section 1 follows the table of contents directly
            self.page_break()
        self._sections += 1
        self._p(None, text, MAROON).style = self._style_id("Heading1")

    def h2(self, text):
        self._p(None, text, MAROON).style = self._style_id("Heading2")

    def _style_id(self, sid):
        return next(s for s in self.d.styles if s.style_id == sid)

    def lead(self, text):
        self._p(None, text).style = self._style_id("FirstParagraph")

    def text(self, text):
        self._p(None, text).style = self._style_id("BodyText")

    def block(self, text):
        self._p(None, text).style = self._style_id("BlockText")

    def items(self, items, numbered=False):
        num = self._new_num() if numbered else BULLET_NUM
        for it in items:
            p = self._p(None, it)
            p.style = self._style_id("Compact")
            p._p.get_or_add_pPr().insert(0, parse_xml(
                f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num}"/></w:numPr>'))

    def _new_num(self):
        """A numbered list that restarts at 1 (as each list in the procedure does)."""
        self.next_num += 1
        base = next(n for n in self.numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == DECIMAL_NUM)
        new = copy.deepcopy(base)
        new.set(qn("w:numId"), str(self.next_num))
        new.append(parse_xml(f'<w:lvlOverride {nsdecls("w")} w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>'))
        self.numbering.append(new)
        return str(self.next_num)

    def code(self, lines):
        p = self._p(None, "")
        p.style = self._style_id("SourceCode")
        for i, line in enumerate(lines):
            if i:
                p.add_run().add_break()
            r = p.add_run(line)
            r._r.get_or_add_rPr().append(parse_xml(f'<w:rStyle {nsdecls("w")} w:val="VerbatimChar"/>'))

    def table(self, head, rows, widths=None):
        cols = len(head)
        widths = widths or [1] * cols
        total = 9360                                                     # 6.5 inch text width, in twips
        grid = "".join(f'<w:gridCol w:w="{int(total * w / sum(widths))}"/>' for w in widths)
        border = "".join(f'<w:{s} w:val="single" w:sz="4" w:color="{GRID}"/>'
                         for s in ("top", "left", "bottom", "right", "insideH", "insideV"))

        white = '<w:color w:val="FFFFFF"/>'

        def cell(text, header=False):
            runs = ""
            for chunk, bold in self._inline(str(text)):
                b = "<w:b/>" if header or bold else ""
                c = white if header else ""
                runs += (f'<w:r><w:rPr>{b}{c}<w:sz w:val="{CELL_SIZE}"/></w:rPr>'
                         f'<w:t xml:space="preserve">{_esc(chunk)}</w:t></w:r>')
            shade = f'<w:shd w:val="clear" w:color="auto" w:fill="{MAROON}"/>' if header else ""
            return (f'<w:tc><w:tcPr>{shade}</w:tcPr><w:p><w:pPr><w:pStyle w:val="Compact"/><w:jc w:val="left"/>'
                    f'</w:pPr>{runs}</w:p></w:tc>')

        xml = (f'<w:tbl {nsdecls("w")}><w:tblPr><w:tblStyle w:val="Table"/>'
               '<w:tblLook w:firstRow="1" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="0" w:val="0020"/>'
               f'<w:jc w:val="start"/><w:tblBorders>{border}</w:tblBorders><w:tblW w:w="5000" w:type="pct"/></w:tblPr>'
               f'<w:tblGrid>{grid}</w:tblGrid>'
               f'<w:tr><w:trPr><w:tblHeader w:val="true"/></w:trPr>{"".join(cell(h, True) for h in head)}</w:tr>'
               + "".join(f'<w:tr>{"".join(cell(c) for c in r)}</w:tr>' for r in rows) + '</w:tbl>')
        self._add(parse_xml(xml))
        self._p(None, "").style = self._style_id("BodyText")   # spacing after the table, as pandoc does

    def page_break(self):
        p = self.d.add_paragraph()
        p.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)

    def save(self):
        cp = self.d.core_properties
        cp.title = "Results-to-Thesis Monitoring Copilot – Phase 1 Pipeline Documentation"
        cp.author = "IFSA IIT Kharagpur"
        cp.subject = "Pipeline documentation and access guide"
        self.d.save(OUT)


def _esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def facts():
    """Current counts from the database and the verification register."""
    conn = sqlite3.connect(db.DEFAULT_DB)
    q = lambda sql: conn.execute(sql).fetchone()[0]
    f = {"docs": q("SELECT count(*) FROM documents"),
         "source_docs": q("SELECT count(*) FROM documents WHERE type != 'analyst_model'"),
         "pages": q("SELECT count(*) FROM pages"), "scanned": q("SELECT count(*) FROM pages WHERE scanned = 1"),
         "scanned_docs": q("SELECT count(DISTINCT document_id) FROM pages WHERE scanned = 1"),
         "figures": q("SELECT count(*) FROM figures"),
         "verified": q("SELECT count(*) FROM figures WHERE status = 'verified'"),
         "manual": q("SELECT count(*) FROM figures WHERE status = 'manually_verified'"),
         "missing": q("SELECT count(*) FROM figures WHERE status = 'missing'"),
         "low": q("SELECT count(*) FROM figures WHERE status = 'low_confidence'"),
         "estimates": q("SELECT count(*) FROM estimates"), "comparisons": q("SELECT count(*) FROM comparisons"),
         "guidance": q("SELECT count(*) FROM statements WHERE type = 'guidance'"),
         "explanations": q("SELECT count(*) FROM statements WHERE type = 'explanation'"),
         "flags": q("SELECT count(*) FROM statements WHERE type = 'flag'"),
         "approved": q("SELECT count(*) FROM reports WHERE status = 'approved'"),
         "audit": q("SELECT count(*) FROM audit_log")}
    sev = dict(conn.execute("SELECT severity, count(*) FROM comparisons WHERE comparator = 'estimate'"
                            " AND severity IS NOT NULL GROUP BY severity").fetchall())
    f["sev"] = sev
    with open(ROOT / "docs" / "manual_verification.csv", encoding="utf-8-sig") as fh:
        mv = list(csv.DictReader(fh))
    f["register"], f["mismatch"] = len(mv), sum(r["match (Y/N)"] == "N" for r in mv)
    f["spot"] = sum(r["spot_check"] == "YES" for r in mv)
    return f


def build():
    f = facts()
    cfg = config.load_all()
    doc = Doc()
    today = dt.date.today().strftime("%-d %B %Y") if sys.platform != "win32" else dt.date.today().strftime("%#d %B %Y")
    doc.title("Results-to-Thesis Monitoring Copilot", "Phase 1 Pipeline Documentation",
              "IFSA IIT Kharagpur — prepared for Ambit Investment Advisors", today)

    # 1 -----------------------------------------------------------------------------------------
    doc.h1("1. Overview")
    doc.lead("The Phase 1 prototype is built and running: a local, rule-based pipeline that turns each company's "
             "quarterly filings into a fully sourced, analyst-approved Change Report, at no cost to Ambit.")
    doc.table(["Item", "Detail"], [
        ["Project", "Results-to-Thesis Monitoring Copilot"],
        ["Client", "Ambit Investment Advisors"],
        ["Team", "IFSA IIT Kharagpur"],
        ["Phase", "1 — Working prototype (built 6–20 October 2026)"],
        ["Coverage", f"{len(cfg.companies)} companies × 2 quarters (Q4 FY26, Q1 FY27) = 10 Change Reports"],
        ["Status", f"{f['approved']} of 10 reports approved; every figure re-checked against its source"],
        ["Extraction", "Rule-based only (no LLM, no paid API)"],
        ["Runs on", "One local machine (Python 3.11 + SQLite)"],
        ["Code", REPO],
    ], [1, 3])
    doc.h2("What the copilot does")
    doc.text("When a company announces quarterly results, an analyst reads the results release, investor presentation "
             "and call transcript, picks out the figures that matter, compares them with their model and with earlier "
             "quarters, and judges what has changed. The copilot does this first pass and produces one Change Report "
             "per company and quarter. For every tracked metric the report shows:")
    doc.items(["**The reported value**, read from the company's own filing.",
               "**The change** against the prior quarter, the year-ago quarter and the analyst's estimate, each graded "
               "Low, Medium or High by size.",
               "**What management said about it** — an exact quote from the earnings call, with speaker and page.",
               "**Exactly where the figure came from** — XBRL tag and context, or PDF page and table row — and whether "
               "it passed the automated checks."])
    doc.text("Each report opens with a headline summary of the largest moves against estimate, for example:")
    doc.block("PAT attributable to owners of ₹5,455 crore was 8.0% above the estimate of ₹5,051 crore (Medium).")
    doc.h2("Guiding principles")
    doc.items(["**Every figure is sourced.** Each number carries its document and page (or XBRL context), period and basis.",
               "**The analyst model is read-only.** The copilot reads estimates and never writes to the file; its hash "
               "is checked before and after every read.",
               "**The analyst stays in control.** Every item is reviewed and nothing is final without explicit approval; "
               "a later change returns the report to Draft.",
               "**Missing means missing.** An undisclosed or unverified figure is shown in Gaps, never guessed, "
               "zero-filled or swapped for another period."])

    # 2 -----------------------------------------------------------------------------------------
    doc.h1("2. Accessing the Prototype")
    doc.lead("Everything needed to run the copilot is in one Git repository; source documents, the database and the "
             "reports are rebuilt locally with five commands.")
    doc.h2("Where things are")
    doc.table(["What", "Where", "Notes"], [
        ["Source code and configuration", REPO, "Public repository, branch main"],
        ["Documentation", "README.md and docs/ in the repository", "This document is docs/Pipeline_Documentation.docx"],
        ["Change Reports (Excel and PDF)", "output/ on the machine that ran the pipeline",
         "Not in Git; regenerated with python run.py export"],
        ["Source documents and database", "data/ on the machine that ran the pipeline",
         "Not in Git; downloaded from NSE and rebuilt locally"],
        ["Review screen", "http://localhost:8501", "Runs locally while streamlit is running"],
    ], [2, 2.5, 2.5])
    doc.h2("Prerequisites")
    doc.items(["Python 3.11 (Windows, macOS or Linux).",
               "Git, to clone the repository.",
               "Internet access to nseindia.com, only for the one-off document download.",
               "About 400 MB of disk space for the source documents."])
    doc.h2("Install")
    doc.code([f"git clone {REPO}.git",
              "cd Ambit-Working-Prototype",
              "py -3.11 -m venv .venv",
              ".venv\\Scripts\\activate",
              "pip install -r requirements.txt"])
    doc.text("On macOS or Linux, create the environment with python3.11 -m venv .venv and activate it with "
             "source .venv/bin/activate.")
    doc.h2("Build the data and the reports")
    doc.code(["python scripts/collect_sources.py",
              "python run.py init-db",
              "python run.py ingest --company all --quarter all",
              "python run.py run    --company all --quarter all",
              "python run.py export --company all --quarter all"])
    doc.items(["Download the filings from NSE into data/ and record each file's URL, date and SHA-256.",
               "Create the SQLite database.",
               "Register every document (hash-checked) and store each PDF page's text.",
               "Extract figures, read estimates, run the checks, compare and grade, and extract transcript statements.",
               "Write every report to Excel and PDF in output/."], numbered=True)
    doc.text("The company option accepts TITAN, ASIANPAINT, MM, INFY, RELIANCE or all; the quarter option accepts "
             "Q4FY25, Q1FY26, Q3FY26, Q4FY26, Q1FY27 or all. A single report can be rebuilt on its own:")
    doc.code(["python run.py run --company MM --quarter Q1FY27"])
    doc.text("This also processes the report's comparator quarters.")
    doc.h2("Open the review screen")
    doc.code(["streamlit run app.py --server.address localhost"])
    doc.items(["Open http://localhost:8501 in a browser.",
               "Enter the reviewer's name in the sidebar and pick a company and quarter.",
               "Open any row to see the source behind the figure: the PDF page image, the XBRL fact and context, or the "
               "analyst-model cell.",
               "Approve, reject, comment on or correct each item; then click Approve report and Export Excel + PDF.",
               "Downloads are offered on the screen; the files are also written to output/."], numbered=True)
    doc.h2("Command reference")
    doc.table(["Command", "Does"], [
        ["python run.py check-config", "Validate every configuration file; stops with a clear error on unknown units, duplicate IDs or undefined inputs"],
        ["python run.py ingest …", "Register documents (hash-checked) and store each PDF page's text"],
        ["python run.py run …", "Extract XBRL and PDF figures, read estimates, run checks, compare and grade, extract transcript statements"],
        ["python run.py export …", "Write each report to Excel and PDF (APPROVED from the saved snapshot, otherwise DRAFT)"],
        ["python scripts/verify_register.py", "Re-read every displayed figure from its source into docs/manual_verification.csv"],
        ["python scripts/check_procedure.py", "Cross-check the build against every requirement of the Phase 1 procedure"],
        ["pytest", "Run the automated test suite"],
    ], [2, 4])

    # 3 -----------------------------------------------------------------------------------------
    doc.h1("3. Pipeline at a Glance")
    doc.lead("The prototype is a single Python application with one SQLite database: a straight pipeline from source "
             "documents to an analyst-approved report, with every step written to an append-only audit log.")
    doc.h2("Pipeline")
    doc.items(["**Collect** — download each company's XBRL filing, results release, investor presentation and call "
               "transcript from NSE and record each file with its URL, date and SHA-256.",
               "**Ingest** — register each document after checking its hash; store the text of every PDF page.",
               "**Extract** — core financials and segments from XBRL; operating KPIs from PDFs with per-company rules; "
               "guidance, explanations and flags from transcripts; estimates from the analyst model (read-only).",
               "**Verify** — automated checks on source, period, basis, unit and reconciliations; failures marked Low confidence.",
               "**Compare** — change vs prior quarter, year-ago quarter and estimate; grade against thresholds.",
               "**Report** — assemble the Change Report for each company and quarter.",
               "**Review** — the analyst checks each item on the review screen and approves the report.",
               "**Export** — the approved report to Excel and PDF from the same saved snapshot."], numbered=True)
    doc.h2("Modules")
    doc.table(["Module", "Responsibility"], [
        ["scripts/collect_sources.py", "Find and download each document from NSE; record it in config/sources.yaml"],
        ["copilot/ingest.py", "Register files, check SHA-256 hashes, store page text, flag scanned pages"],
        ["copilot/xbrl_extract.py", "Parse XBRL contexts, units and facts; map them to metrics and segments"],
        ["copilot/pdf_extract.py", "Apply per-company rules to presentations and fact sheets"],
        ["copilot/estimates.py", "Read mapped cells from the analyst model, read-only"],
        ["copilot/commentary.py", "Speaker turns, guidance, explanations, flags, prior-quarter comparison"],
        ["copilot/verify.py", "Automated checks and confidence status"],
        ["copilot/compare.py", "QoQ, YoY and vs-estimate changes; severity grades"],
        ["copilot/report.py", "Build the report; review state, approval rule, fingerprint"],
        ["copilot/review.py", "Analyst actions: approve, reject, comment, correct"],
        ["copilot/export.py", "Write Excel and PDF from one snapshot"],
        ["app.py / run.py", "Streamlit review screen / command line"],
    ], [2, 4])
    doc.h2("Tech stack")
    doc.table(["Tool", "Used for"], [
        ["Python 3.11", "Application"], ["lxml", "XBRL parsing"],
        ["PyMuPDF", "PDF page text and word positions"], ["SQLite", "Storage and audit log"],
        ["openpyxl", "Read the analyst model; write the Excel report"], ["Streamlit", "Review screen"],
        ["ReportLab", "PDF report"], ["PyYAML", "Configuration and metric dictionary"], ["pytest", "Tests"],
    ], [1, 3])
    doc.text("All libraries are free and open source; no API key is needed.")

    # 4 -----------------------------------------------------------------------------------------
    doc.h1("4. Scope & Coverage")
    doc.lead("Phase 1 covers five listed companies across five sectors, two report quarters and three comparator "
             "quarters, on a consolidated basis.")
    doc.h2("Companies")
    doc.table(["Company", "Sector", "What it tests", "Segments"], [
        [c["name"], c.get("sector", ""), c.get("tests", ""), ", ".join(c["segments"]) or "single segment"]
        for c in cfg.companies.values()], [2, 1.3, 2.5, 2.7])
    doc.h2("Quarters")
    doc.table(["Report quarter", "Prior quarter", "Year-ago quarter", "Estimate"], [
        ["Q4 FY26 (Jan–Mar 2026)", "Q3 FY26 (Oct–Dec 2025)", "Q4 FY25 (Jan–Mar 2025)", "Analyst model, Q4 FY26"],
        ["Q1 FY27 (Apr–Jun 2026)", "Q4 FY26 (Jan–Mar 2026)", "Q1 FY26 (Apr–Jun 2025)", "Analyst model, Q1 FY27"]])
    doc.h2("Source documents")
    doc.table(["Document", "Quarters", "Used for"], [
        ["XBRL results filing (SEBI integrated filing)", "All 5", "Core financials, segments, net debt"],
        ["Results release (board outcome PDF)", "All 5", "Cross-check of revenue and PAT"],
        ["Investor presentation", "Report quarters (Titan: all 5)", "Operating KPIs"],
        ["Earnings-call transcript", "Q3 FY26, Q4 FY26, Q1 FY27", "Guidance, explanations, flags"],
        ["Results press release (Reliance)", "All 5", "Second cross-check source"],
        ["Analyst model (sample)", "Q4 FY26E, Q1 FY27E", "Estimates"],
    ], [2.5, 2, 2])
    doc.text(f"{f['source_docs']} source documents are registered with their SHA-256 hashes ({f['pages']:,} PDF pages "
             f"stored; {f['scanned']} pages in {f['scanned_docs']} documents flagged as scanned). Documents that do "
             "not exist — Infosys files no investor presentation on NSE; Titan files no results press release — are "
             "recorded as missing with the reason.")

    # 5 -----------------------------------------------------------------------------------------
    doc.h1("5. Data Collection & Extraction")
    doc.lead("Figures come from three kinds of source, each read by its own rule-based extractor, and every figure is "
             "stored with exactly where it was found.")
    doc.h2("Core financials from XBRL")
    doc.items(["Each SEBI integrated filing covers only its own quarter, so every quarter's figures come from that "
               "quarter's own filing (all five quarters are downloaded).",
               "A fact is used only if its context covers exactly that quarter's three months (or the quarter-end date "
               "for balance-sheet items) and the filing's basis matches the configured basis.",
               "Values are converted from rupees to exact ₹ crore decimals; the original value and unit are kept.",
               "Segments are matched by the label the company gives them, because their position changes between "
               "filings; an unknown label stops the run.",
               "EBITDA and EBITDA margin are derived, with the formula and input figures recorded."])
    doc.h2("Operating KPIs from PDFs")
    doc.text("Each KPI is located by word position on the page, never as a number found anywhere on the page. Rules "
             "live in config/pdf_rules/<company>.yaml:")
    doc.table(["Rule kind", "Finds", "Used for"], [
        ["grid", "The value where a row label meets a column header (quarter labels or dates)", "Infosys fact sheet, Jio KPI table"],
        ["chart", "The value directly above each quarter label on a chart axis", "Asian Paints, M&M"],
        ["labelled", "The value directly above a caption", "Titan secondary-growth table"],
        ["text", "A sentence pattern for the document's own quarter", "Infosys constant-currency growth"],
    ], [1, 3, 2])
    doc.h2("Transcripts")
    doc.items(["Each call is split into speaker turns using per-company speaker formats; management is identified from "
               "each transcript's own management list.",
               "**Guidance:** a management sentence with a forward-looking phrase and a tracked topic.",
               "**Explanations:** a management sentence linking a topic to a cause (\"due to\", \"driven by\").",
               "**Flags:** auditor, restatement, related-party, regulatory or accounting-policy phrases, skipped when "
               "negated (\"no qualified opinion\"); always marked for analyst review.",
               "Each topic's guidance is placed beside the prior quarter's: repeated / updated, new, or not repeated "
               "(never \"withdrawn\")."])
    doc.h2("Estimates")
    doc.text("Estimates are read from the analyst model through config/model_map.yaml, with openpyxl in read-only mode "
             "and cached values only; the file is never saved and its hash must be unchanged after the read. No real "
             "model was available, so a sample model (models/sample_analyst_model.xlsx) holds synthetic estimates; "
             "every report labels them SAMPLE.")

    # 6 -----------------------------------------------------------------------------------------
    doc.h1("6. Verification, Comparison & Grading")
    doc.lead("Accuracy rests on three layers: automated checks on every figure, an independent re-read of every "
             "figure against its source, and analyst approval of every report.")
    doc.h2("Automated checks")
    doc.table(["Check", "How", "On failure"], [
        ["Source recorded", "Document and page, or XBRL element and context", "Missing (listed in Gaps)"],
        ["Value on cited page", "PDF figure: the printed value is on the cited page", "Low confidence"],
        ["Period and basis", "Exact three-month context; basis matches the configuration", "Low confidence"],
        ["Unit", "XBRL unit is INR (money) or INR per share (EPS)", "Low confidence"],
        ["Cross-source", "XBRL revenue and PAT appear in that quarter's results release or press release, or a later release's comparator column", "Low confidence"],
        ["Segment sum", "Segment revenues − inter-segment = segment total = revenue", "Low confidence"],
        ["Derived inputs", "Every input of a derived figure is verified", "Low confidence"],
    ], [1.5, 3.5, 1.5])
    doc.text("A Low-confidence figure stays in the metric table, clearly marked, and is excluded from the headline "
             "summary until the analyst resolves it by checking it against the source.")
    doc.h2("Comparisons")
    doc.items(["Amounts: % change. Margins and other % levels: basis points. Growth rates disclosed by the company: "
               "percentage points.",
               "Zero or negative base: the absolute change only, labelled n.m. (not meaningful).",
               "Comparator copy: the value read from the report quarter's own document first, then the comparator "
               "quarter's own document.",
               "A missing comparator or estimate is recorded as a gap with its reason."])
    doc.h2("Severity bands")
    doc.text("Severity grades the size of a move, not whether it is good or bad. Moves against the estimate use the "
             "main bands; moves against the prior and year-ago quarters use looser bands so that ordinary growth does "
             "not grade High.")
    doc.table(["Group", "vs estimate: Low / Medium / High", "vs prior and year-ago"], [
        ["Revenue, segment revenue", "< 2% / 2–5% / > 5%", "< 5% / 5–15% / > 15%"],
        ["EBITDA, segment result, PAT, EPS, net debt", "< 5% / 5–10% / > 10%", "< 10% / 10–25% / > 25%"],
        ["Margins and % levels", "< 50 / 50–150 / > 150 bps", "< 100 / 100–300 / > 300 bps"],
        ["Volume and count KPIs", "< 5% / 5–10% / > 10%", "< 10% / 10–20% / > 20%"],
        ["Growth-rate KPIs", "< 2 / 2–5 / > 5 pp", "< 5 / 5–10 / > 10 pp"],
    ], [2.5, 2, 2])

    # 7 -----------------------------------------------------------------------------------------
    doc.h1("7. Change Report, Review & Export")
    doc.lead("One report per company per quarter, built from the database, reviewed on screen and exported to Excel "
             "and PDF from the same approved snapshot.")
    doc.h2("Report sections")
    doc.table(["Section", "Contents"], [
        ["Header", "Company, quarter, basis, status (Draft / Approved), approver and time, documents used, company notes"],
        ["Headline summary", "The 3–5 largest moves against estimate, from verified figures only, as plain sentences"],
        ["Metric table", "Reported value; prior quarter, year-ago and estimate with change; severity; management explanation; source; verification status"],
        ["Guidance changes", "Per topic, up to two current and two prior-quarter quotes with page references"],
        ["Governance / accounting flags", "Each flag quote with speaker, page and \"for analyst review\""],
        ["Gaps", "Figures not disclosed, comparisons without a comparator, and unverified figures, with the reason"],
        ["Sources", "Every document used, with its publication date and SHA-256"],
    ], [1.5, 4.5])
    doc.h2("Approval rule")
    doc.items(["Every item — each metric row, guidance candidate and flag — must be approved, rejected or corrected.",
               "A Low-confidence figure can be approved only with a note saying what was checked against the source; "
               "it then becomes Manually verified, and the check survives re-runs while the value is unchanged.",
               "A correction needs a reason; the original extraction is kept and the comparisons are recomputed.",
               "Approval records the reviewer and time and saves a snapshot with a fingerprint of its content. If any "
               "input changes later, the report returns to Draft. Exporting never counts as approving."])
    doc.h2("Exports")
    doc.text("Excel (sheets: Report, Guidance, Flags, Gaps, Sources, Audit log) and PDF (landscape A4) are written "
             "from the same display rows, so they always match. Unapproved reports carry DRAFT on every sheet and a "
             "watermark on every page; files are named <COMPANY>_<QUARTER>_<DRAFT|APPROVED>.")

    # 8 -----------------------------------------------------------------------------------------
    doc.h1("8. Configuration & Data Model")
    doc.lead("All business rules live in YAML files that an analyst can read, and all results live in one SQLite "
             "database; the code holds no company-specific numbers.")
    doc.h2("Configuration files")
    doc.table(["File", "Holds"], [
        ["companies.yaml", "Company ID, name, NSE/BSE code, sector, basis, segment aliases, report notes"],
        ["metrics.yaml", "The metric dictionary (see docs/metric_dictionary.md)"],
        ["thresholds.yaml", "Low / Medium / High bands per metric group, vs estimate and vs prior / year-ago"],
        ["sources.yaml", "Every document with type, file path, period, URL, publication date and SHA-256"],
        ["pdf_rules/<company>.yaml", "Page anchors, row labels, column headers or chart regions for each KPI"],
        ["commentary_rules.yaml", "Speaker formats, management rosters, guidance / explanation / flag phrases, topics"],
        ["model_map.yaml", "Analyst-model sheet and cell for each metric and quarter, units, estimate date"],
    ], [2, 4])
    doc.h2("Database tables")
    doc.table(["Table", "Holds"], [
        ["documents", "Every source document with hash, URL and dates"],
        ["pages", "The text of every PDF page, with a scanned flag"],
        ["figures", "Every value with period, basis, segment, unit, source (page or XBRL context), method and status"],
        ["estimates", "Every estimate with its sheet and cell"],
        ["comparisons", "Each change vs prior, year-ago and estimate, with severity or gap reason"],
        ["statements", "Guidance, explanations and flags as exact quotes with page and speaker"],
        ["reviews", "Every analyst action with reviewer, comment and old / new value"],
        ["reports", "Report status, approved snapshot, fingerprint, approver and time"],
        ["audit_log", "Append-only record of every ingest, extraction, check, flag, review, approval and export"],
    ], [1.5, 4.5])
    doc.h2("Data rules")
    doc.items(["Values are stored as exact decimals, never floats.",
               "Monetary figures are normalised to ₹ crore; the original value and unit are kept beside them.",
               "A blank, \"–\" or nil value is missing, never zero.",
               "Figure status is Verified, Low confidence, Missing, or Manually verified once checked against the source.",
               "An analyst correction is saved as a new value with a reason; the original extraction is kept.",
               "The audit log is append-only: updates and deletes are rejected by the database."])

    # 9 -----------------------------------------------------------------------------------------
    sev = f["sev"]
    doc.h1("9. Results & Quality Assurance")
    doc.lead(f"All 10 Change Reports are built and {f['approved']} are approved; every one of the {f['register']} "
             f"figures they display was re-read from its source with {f['mismatch']} mismatches.")
    doc.h2("Phase 1 results")
    doc.table(["Measure", "Result"], [
        ["Documents registered", f"{f['docs']} ({f['source_docs']} source documents + analyst model), {f['pages']:,} PDF pages"],
        ["Figures extracted", f"{f['figures']} — {f['verified']} verified, {f['manual']} manually verified, "
                              f"{f['missing']} missing with a reason, {f['low']} low confidence"],
        ["Estimates read", f"{f['estimates']}, model hash unchanged on every read"],
        ["Comparisons", f"{f['comparisons']}; vs estimate: {sev.get('low', 0)} Low, {sev.get('medium', 0)} Medium, "
                        f"{sev.get('high', 0)} High"],
        ["Transcript statements", f"{f['guidance']} guidance, {f['explanations']} explanations, {f['flags']} flags"],
        ["Change Reports", f"10 built, {f['approved']} approved, exported to Excel and PDF"],
        ["Audit log", f"{f['audit']:,} entries from ingest to export"],
    ], [2, 4])
    doc.h2("Quality assurance")
    doc.table(["Check", "Result"], [
        ["Automated tests (pytest)", "Number parsing, XBRL selection, PDF rules, comparisons, severity boundaries, "
                                     "analyst model, commentary, approval flow — all passing"],
        ["Independent re-read of every figure", f"{f['register']} figures re-read through a separate code path "
                                                f"(raw XML, fresh PDF page, recomputation, model cell): "
                                                f"{f['mismatch']} mismatches; {f['spot']} rows flagged for a hand spot-check"],
        ["Procedure cross-check", "Every requirement of the Phase 1 procedure tested against the build: 44 pass, "
                                  "8 documented departures, 0 fail (docs/procedure_crosscheck.md)"],
        ["Clean rebuild", "A clean copy rebuilt from the README commands gives identical results"],
    ], [2, 4])

    # 10 ----------------------------------------------------------------------------------------
    doc.h1("10. Decisions & Limitations")
    doc.lead("Phase 1 departs from the original procedure where the real filings required it; each decision is "
             "recorded in the repository's docs/ folder.")
    doc.h2("Decisions that changed the plan")
    doc.table(["Topic", "Procedure", "As built"], [
        ["Sample", "3 manufacturing companies", "5 companies in 5 sectors (Titan, Asian Paints, M&M, Infosys, Reliance)"],
        ["Sign-off", "Ambit approves scope", "Scope decided by the team for the prototype"],
        ["EBITDA", "PBEIT + finance costs + D&A − other income", "Also excludes the share of associates and JVs"],
        ["Segments", "From the results release PDF", "From XBRL (structured, same filing)"],
        ["KPI reading", "pdfplumber tables", "PyMuPDF word positions (most KPIs are charts)"],
        ["Growth-rate KPIs", "% change", "Percentage points, on their own bands"],
        ["QoQ / YoY grades", "Same bands as vs estimate", "Looser bands, so ordinary growth does not grade High"],
        ["Net debt", "Where disclosed", "Q4 only; not shown for M&M (NBFC lending book) or Titan (gold metal loans)"],
        ["Estimates", "Analyst model", "Synthetic sample model, labelled SAMPLE"],
    ], [1.3, 2, 3])
    doc.h2("Known limitations")
    doc.items(["PDF rules are written per company and per layout; a changed layout needs a rule update.",
               "Scanned (image-only) PDFs are not read automatically; Titan's Q1 FY26 deck is one, so those year-ago "
               "KPIs are gaps.",
               "Transcript rules find keyword candidates; the analyst judges their meaning.",
               "Operating cash flow and capex are filed only for the full year, and net debt only at 31 March.",
               "Correcting an input figure does not recompute the figures derived from it.",
               "Reviewer identity is a name entered on the local screen, not a secure login."])
    doc.h2("Open items")
    doc.items(["Record the 2–3-minute demo video.",
               f"Hand spot-check of the {f['spot']} flagged rows in docs/manual_verification.csv by a team member.",
               "Item-by-item analyst review of the reports (the current approvals were a bulk approval on instruction).",
               "Phase 2: map Ambit's real analyst model; add quarterly cash flow and balance-sheet coverage."])
    doc.save()
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
