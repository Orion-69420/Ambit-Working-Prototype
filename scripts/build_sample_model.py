"""Build the SAMPLE analyst model (Step 7) and its cell map.

    python scripts/build_sample_model.py

Writes models/sample_analyst_model.xlsx and config/model_map.yaml. Estimates are synthetic:
each is the reported figure moved by a fixed offset chosen from the metric, segment and quarter
(so the result is reproducible and gives a mix of Low/Medium/High variances). EBITDA margin is
recomputed from the estimated EBITDA and revenue, and EPS moves with PAT, as in a real model.

Run once after Steps 5-6; the workbook is committed so everyone reads the same file. Rebuilding
changes its hash (xlsx files embed timestamps).
"""
import datetime as dt
import hashlib
import pathlib
import sqlite3
import sys
from decimal import ROUND_HALF_UP, Decimal

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from copilot import config, db  # noqa: E402

WORKBOOK = ROOT / "models" / "sample_analyst_model.xlsx"
MODEL_MAP = ROOT / "config" / "model_map.yaml"
QUARTERS = list(config.TARGET_QUARTERS)
ESTIMATE_DATES = {"Q4FY26": "2026-04-10", "Q1FY27": "2026-07-10"}   # before any company reported

# variance of actual vs estimate that the sample aims for
OFFSETS = {
    "pct": [-0.12, -0.06, -0.035, -0.015, -0.005, 0.005, 0.015, 0.03, 0.045, 0.08],
    "bps": [-180, -100, -60, -30, -10, 10, 30, 60, 100, 180],
    "pp": [-7, -4, -2.5, -1, -0.5, 0.5, 1, 2.5, 4, 7],
}
PLACES = {"inr_crore": "1", "inr_per_share": "0.01", "percent": "0.1", "pct_growth": "0.1",
          "thousand_units": "0.1", "million": "0.1", "inr_per_month": "0.1", "gb_per_month": "0.1"}
SKIP = {"operating_cash_flow", "capex"}          # annual-only: never compared


def offset(kind, *key):
    h = int(hashlib.sha256("|".join(map(str, key)).encode()).hexdigest()[:8], 16)
    return OFFSETS[kind][h % len(OFFSETS[kind])]


def estimate(actual, kind, off):
    if kind == "pct":
        return actual / (1 + Decimal(str(off)))
    if kind == "bps":
        return actual - Decimal(str(off)) / 100
    return actual - Decimal(str(off))


def rnd(v, unit):
    return v.quantize(Decimal(PLACES[unit]), rounding=ROUND_HALF_UP)


def actual(conn, company, metric, segment, quarter):
    """Reported value for the quarter: XBRL from its own filing, KPIs from any document showing it."""
    row = conn.execute(
        "SELECT value FROM figures WHERE company = ? AND metric = ? AND period = ? AND value IS NOT NULL"
        " AND segment IS ? ORDER BY (quarter = period) DESC, quarter DESC LIMIT 1",
        (company, metric, quarter, segment)).fetchone()
    return db.from_db(row[0]) if row else None


def rows_for(conn, cfg, company):
    """[(metric, segment, label, unit, {quarter: estimate})] in report order."""
    comp = cfg.companies[company]
    out = []
    for m in cfg.metrics_for(company):
        if m["id"] in SKIP:
            continue
        segments = list(comp["segments"]) if m.get("per_segment") else [None]
        for seg in segments:
            ests = {}
            for q in QUARTERS:
                if m["id"] == "net_debt" and q != "Q4FY26":
                    continue                                    # balance sheet only at 31 March
                a = actual(conn, company, m["id"], seg, q)
                if a is None or m["type"] == "derived" and m["id"] == "ebitda_margin":
                    continue
                if m["id"] == "basic_eps":
                    continue                                    # set from PAT below
                ests[q] = rnd(estimate(a, m["change_as"], offset(m["change_as"], company, m["id"], seg, q)), m["unit"])
            out.append([m, seg, ests])
    by = {(r[0]["id"], r[1]): r[2] for r in out}
    for r in out:
        m, seg, ests = r
        for q in QUARTERS:
            if m["id"] == "ebitda_margin" and q in by.get(("ebitda", None), {}) and q in by.get(("revenue", None), {}):
                ests[q] = rnd(by[("ebitda", None)][q] / by[("revenue", None)][q] * 100, "percent")
            if m["id"] == "basic_eps":
                pat_a, eps_a = actual(conn, company, "pat_owners", None, q), actual(conn, company, "basic_eps", None, q)
                pat_e = by.get(("pat_owners", None), {}).get(q)
                if pat_a and eps_a and pat_e:
                    ests[q] = rnd(eps_a * pat_e / pat_a, "inr_per_share")
    return [(m, seg, ests) for m, seg, ests in out if ests]


def main():
    cfg = config.load_all()
    conn = sqlite3.connect(db.DEFAULT_DB)
    wb = Workbook()
    about = wb.active
    about.title = "About"
    for i, line in enumerate([
        "SAMPLE ANALYST MODEL - SYNTHETIC ESTIMATES",
        "Not an analyst view. Built for the Phase 1 prototype because no real model was available.",
        "Each estimate is the reported figure moved by a fixed offset (see scripts/build_sample_model.py).",
        "EBITDA margin = EBITDA estimate / revenue estimate; EPS moves with PAT.",
        "Money in INR crore; percentages in %; growth rates in % YoY.",
        f"Built {dt.date.today().isoformat()} from the extracted figures (Steps 5-6).",
    ], start=1):
        about.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), color="C00000" if i <= 2 else None)

    model_map = {"workbook": WORKBOOK.relative_to(ROOT).as_posix(), "sample": True,
                 "note": "Generated by scripts/build_sample_model.py; cells are read-only inputs", "companies": {}}
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for company in cfg.companies:
        ws = wb.create_sheet(company)
        ws["A1"] = f"{cfg.companies[company]['name']} - SAMPLE ESTIMATES (synthetic)"
        ws["A1"].font = Font(bold=True, color="C00000")
        heads = ["Metric", "Segment", "Unit", *[f"{q}E" for q in QUARTERS], "metric id"]
        for c, h in enumerate(heads, start=1):
            cell = ws.cell(row=3, column=c, value=h)
            cell.font, cell.fill = Font(bold=True), header_fill
        qcol = {q: get_column_letter(4 + i) for i, q in enumerate(QUARTERS)}
        ws.cell(row=4, column=1, value="Estimate date")
        for q, col in qcol.items():
            ws[f"{col}4"] = ESTIMATE_DATES[q]
        entry = {"sheet": company, "estimate_date": {q: f"{col}4" for q, col in qcol.items()}, "cells": []}
        r = 5
        for m, seg, ests in rows_for(conn, cfg, company):
            ws.cell(row=r, column=1, value=m["name"])
            ws.cell(row=r, column=2, value=seg)
            ws.cell(row=r, column=3, value=cfg.units[m["unit"]]["label"])
            ws.cell(row=r, column=4 + len(QUARTERS), value=m["id"])
            cells = {"metric": m["id"], "segment": seg, "unit": m["unit"]}
            for q, col in qcol.items():
                if q in ests:
                    ws[f"{col}{r}"] = ests[q]
                    cells[q] = f"{col}{r}"
            entry["cells"].append(cells)
            r += 1
        for col, width in zip("ABCDEF", (38, 34, 12, 12, 12, 24)):
            ws.column_dimensions[col].width = width
        model_map["companies"][company] = entry

    wb.properties.creator = "IFSA IIT Kharagpur - Phase 1 prototype (sample model)"
    WORKBOOK.parent.mkdir(exist_ok=True)
    wb.save(WORKBOOK)
    header = ("# Analyst model cell map (Step 7), generated by scripts/build_sample_model.py.\n"
              "# Each entry: metric (and segment) -> cell per quarter on the company's sheet.\n")
    MODEL_MAP.write_text(header + yaml.safe_dump(model_map, sort_keys=False, allow_unicode=True), encoding="utf8")
    n = sum(len([k for k in c if k in QUARTERS]) for e in model_map["companies"].values() for c in e["cells"])
    print(f"wrote {WORKBOOK.relative_to(ROOT)} ({n} estimates) and {MODEL_MAP.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
